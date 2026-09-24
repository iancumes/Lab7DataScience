"""Obtain the five official INE Personas workbooks and their dictionaries."""

from __future__ import annotations

import concurrent.futures
import hashlib
import json
from datetime import datetime, timezone
from pathlib import Path
from zipfile import ZipFile

import requests
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
BASE = "https://www.ine.gob.gt/wp-content/uploads"
FILES = [
    ("2025T1", "2026/01/Personas_ENEIC_T1_2025.xlsx", "2025/11/Diccionario_Personas_ENEIC_I-2025.xlsx", 51588, 270),
    ("2025T2", "2026/01/Personas-ENEIC-T2-2025.xlsx", "2025/11/Diccionario_Personas_ENEIC_II-2025.xlsx", 51167, 270),
    ("2025T3", "2026/05/Base-de-datos-Personas-ENEIC-III-2025.xlsx", "2026/05/Diccionario-Personas-ENEIC-III-2025.xlsx", 51583, 270),
    ("2025T4", "2026/06/Base-de-datos-Personas-ENEIC-IV-2025.xlsx", "2026/06/Diccionario-Personas-ENEIC-IV-2025.xlsx", 49338, 302),
    ("2026T1", "2026/09/Base-de-datos-Personas-ENEIC-I-2026.xlsx", "2026/09/Diccionario-Personas-ENEIC-I-2026.xlsx", 49843, 270),
]


def obtain(period: str, kind: str, rel_url: str) -> dict:
    directory = ROOT / "data" / ("raw" if kind == "personas" else "dictionaries")
    directory.mkdir(parents=True, exist_ok=True)
    path = directory / f"{period}_{kind}.xlsx"
    url = f"{BASE}/{rel_url}"
    if not path.exists():
        temporary = path.with_suffix(".download")
        try:
            with requests.get(url, stream=True, timeout=(20, 180)) as response:
                response.raise_for_status()
                with temporary.open("wb") as target:
                    for block in response.iter_content(chunk_size=1024 * 1024):
                        if block:
                            target.write(block)
            with ZipFile(temporary) as workbook:
                if "xl/workbook.xml" not in workbook.namelist():
                    raise ValueError("La respuesta no es un libro Excel")
            temporary.replace(path)
        finally:
            temporary.unlink(missing_ok=True)
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for block in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(block)
    with ZipFile(path) as workbook:
        assert "xl/workbook.xml" in workbook.namelist(), path
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows = sheet.max_row - 1 if kind == "personas" else sheet.max_row
        columns = sheet.max_column
        sheet_name = sheet.title
        header = (" ".join(str(value) for row in sheet.iter_rows(min_row=1, max_row=2, values_only=True)
                           for value in row if value is not None)[:180] if kind == "diccionario" else None)
    finally:
        workbook.close()
    return {
        "periodo": period,
        "tipo": kind,
        "archivo": path.relative_to(ROOT).as_posix(),
        "url": url,
        "fecha_descarga_utc": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "sha256": digest.hexdigest(),
        "bytes": path.stat().st_size,
        "filas_observadas": rows,
        "columnas_observadas": columns,
        "hoja": sheet_name,
        **({"encabezado_inicial": header} if header is not None else {}),
    }


def main() -> None:
    existing_path = ROOT / "sources_manifest.json"
    existing = json.loads(existing_path.read_text(encoding="utf-8")) if existing_path.exists() else {}
    existing_files = {(x["periodo"], x["tipo"]): x for x in existing.get("archivos", [])}
    jobs = [(period, "personas", data) for period, data, _, _, _ in FILES]
    jobs += [(period, "diccionario", dictionary) for period, _, dictionary, _, _ in FILES]
    with concurrent.futures.ThreadPoolExecutor(max_workers=3) as pool:
        futures = {pool.submit(obtain, *job): job for job in jobs}
        records = []
        for future in concurrent.futures.as_completed(futures):
            job = futures[future]
            try:
                record = future.result()
            except Exception as error:
                raise RuntimeError(f"No se pudo obtener {job[0]} {job[1]} de {BASE}/{job[2]}: {error}") from error
            previous = existing_files.get((record["periodo"], record["tipo"]))
            if previous and previous["sha256"] == record["sha256"]:
                record["fecha_descarga_utc"] = previous["fecha_descarga_utc"]
            records.append(record)
            print(f'{record["periodo"]} {record["tipo"]}: {record["bytes"]:,} bytes, {record["sha256"][:12]}', flush=True)
    records.sort(key=lambda r: (r["periodo"], r["tipo"]))
    manifest = {
        "fuente": "https://www.ine.gob.gt/encuesta-nacional-de-empleo-e-ingresos/",
        "nota": "El enlace de diccionario 2026T1 figura como 2023 en el texto de la página; verificar contenido.",
        "esperado": {p: {"filas": rows, "columnas": cols} for p, _, _, rows, cols in FILES},
        "archivos": records,
    }
    existing_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


if __name__ == "__main__":
    main()
